# -*- coding: utf-8 -*-
"""fetch/hood.py 期次表头（_periods）的护栏测试 —— 离线，夹具是现编的 xlsx。

跑法: python3 fetch/test_hood.py

━━ 这套测试守的是哪一件事 ━━
2026-09-11 起 hood 天天 FAIL：`DateParseError: Unknown datetime string format, unable to
parse: NONE-07`。August 2026 Monthly Metrics 把左边那个年份块（Jul..Dec 2025）的年份格
留空了：D6:I6 仍合并着，但 D6 没有值，整个第 6 行只剩 K6=2026。旧 _periods 从左往右
沿用年份，D7「Jul」拿到 None，拼出 'None-07' 交给 pandas。结果是 8 月整月进不了 series，
报错里却没有文件名、sheet 名、列号。

修法见 _periods 的 docstring：表头逐期连续 + 年份格定锚；读不到年份的格子按日历推，
但必须与标题「Monthly Metrics Report for August 2026」互证一致才放行。本文件四组：
  · TestAug2026Layout —— 正题：09-10 那份的版式要读对；推不稳时要抛得具体。
  · TestControlGroup —— 对照组：年份格齐全的老版式结果不变、不出声；标题没改也不拦。
  · TestLoudFailures —— 旧写法下「照读不误但读错」的几种坏法，现在都要抛。
  · TestRealCachedWorkbook —— 夹具对得上真文件：拿 cache/ 里那份 8 月工作簿实跑。
    cache 不进 git，没有这份文件就 skip —— skip 不等于通过。

━━ 夹具坐标照抄真文件 ━━
第 3 行标题、第 6 行年份、第 7 行月份；8 月版 D..I 与 K..R 两块、J 列空隔；7 月版 D..J 与
L..R 两块、K 列空隔；T/U 是 M/M、Y/Y。都来自 2026-09-12 对 cache 里 July / August 2026
两份 Monthly Metrics 的实读（openpyxl 与解包原始 XML 两种读法结论一致）。
指标行只求让 parse_sheet 找齐 25 列，数值是编的：第 r 行第 j 个数据列填 r*100+j，
读出来的数一眼看得出来自哪一格。

━━ 修复前 vs 修复后（2026-09-12 实测，同一份测试文件分别打两版 hood.py）━━
    修复前（HEAD 7d67e7f）：Ran 11 tests，FAILED (failures=2, errors=7)
        errors 7：TestAug2026Layout 四条、无年份格、真文件 —— 六条全是 DateParseError NONE-07；
                  季度表那条是 TypeError（旧 _periods 没有 where、report 两个参数）
        failures 2：月份断列、年份格矛盾 —— 旧写法照读不误，RuntimeError not raised
    修复后：Ran 11 tests，OK
两版都绿的 TestControlGroup 是对照组 —— 它证明年份格齐全的文件结果一个字都没动。
"""

import contextlib
import io
import os
import sys
import tempfile
import unittest

import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

import hood      # noqa: E402

MON = hood.MON

# ── 表头夹具：列字母 → 单元格内容，全部照抄真文件 ──
AUG_TITLE = 'Monthly Metrics Report for August 2026'
AUG_YEARS = {'K': 2026}                                  # D6 空 —— 这就是 09-10 那份
AUG_HEADS = {**dict(zip('DEFGHI', MON[6:])), **dict(zip('KLMNOPQR', MON[:8]))}
AUG_MERGES = ('D6:I6', 'K6:Q6', 'T6:U6')

JUL_TITLE = 'Monthly Metrics Report for July 2026'
JUL_YEARS = {'D': 2025, 'L': 2026}
JUL_HEADS = {**dict(zip('DEFGHIJ', MON[5:])), **dict(zip('LMNOPQR', MON[:7]))}
JUL_MERGES = ('D6:J6', 'L6:R6', 'T6:U6')

# spec 里的本名（不含别名）按字面顺序排，章节标题直接用 spec 的章节键 —— _section_of 是前缀匹配
PRIMARY = [(k, v) for k, v in hood.MONTHLY_SPEC.items() if k not in hood.MONTHLY_ALIASES]

REAL_AUG = os.path.join(os.path.dirname(HERE), 'cache',
                        'hood_2026-08_Robinhood_Markets_Inc_August_2026_Monthly_Metrics_2099b7d0_x.xlsx')


def make_book(path, title, years, heads, merges=()):
    """照真文件坐标摆一张 'Monthly Metrics'，返回 {列名: 所在行号}。"""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Monthly Metrics'
    ws['B2'] = 'Robinhood Markets, Inc. and Consolidated Subsidiaries'
    if title:
        ws['B3'] = title
    ws['B4'] = '(Unaudited)'
    for col, v in years.items():
        ws[f'{col}6'] = v
    ws['T6'] = 'Change'
    for col, v in heads.items():
        ws[f'{col}7'] = v
    ws['T7'], ws['U7'] = 'M/M', 'Y/Y'
    for rng in merges:
        ws.merge_cells(rng)
    data_cols = sorted(heads, key=openpyxl.utils.column_index_from_string)
    r, section, where = 9, None, {}
    for (sec, lab), col in PRIMARY:
        if sec != section:
            ws[f'B{r}'] = sec
            section, r = sec, r + 1
        ws[f'B{r}'] = lab
        for j, c in enumerate(data_cols):
            ws[f'{c}{r}'] = r * 100 + j
        where[col] = r
        r += 1
    wb.save(path)
    return where


def parse(path):
    """parse_workbook，同时收下它打的日志（推算年份时会打一行）。"""
    buf = io.StringIO()
    with contextlib.redirect_stdout(buf):
        m, _q = hood.parse_workbook(path)
    return m, buf.getvalue()


def months(first, last):
    return [str(p) for p in pd.period_range(first, last, freq='M')]


class _Books(unittest.TestCase):
    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()

    def tearDown(self):
        self._tmp.cleanup()

    def book(self, name, **kw):
        path = os.path.join(self._tmp.name, name)
        return path, make_book(path, **kw)


class TestAug2026Layout(_Books):
    """正题：D6 空、K6=2026、标题 August 2026。"""

    def test_blank_left_year_block_reads_as_2025(self):
        """**修复前红、修复后绿的那条。** 左块推成 2025，R 列落在 2026-08，并且日志里留痕。"""
        path, rows = self.book('aug.xlsx', title=AUG_TITLE, years=AUG_YEARS,
                               heads=AUG_HEADS, merges=AUG_MERGES)
        m, log = parse(path)
        self.assertEqual([str(p) for p in m.index], months('2025-07', '2026-08'))
        r = rows['funded_customers_mn']
        self.assertEqual(m.loc[pd.Period('2025-07', 'M'), 'funded_customers_mn'], r * 100 + 0)   # D 列
        self.assertEqual(m.loc[pd.Period('2026-08', 'M'), 'funded_customers_mn'], r * 100 + 13)  # R 列
        # 推算必须看得见：哪几列、推成了哪几个月、拿什么互证的
        for token in ('aug.xlsx「Monthly Metrics」', 'D..I', '2025-07..2025-12', 'K6=2026', '2026-08'):
            self.assertIn(token, log)
        # 右块 K..R 头上有年份格，不该被算作「推出来的」
        self.assertNotIn('K..R', log)

    def test_no_title_means_no_guess_and_error_is_specific(self):
        """同一版式但没有标题 ⇒ 不推，抛 RuntimeError；报错要指得出文件、sheet、列、锚点，
        不能再是一句 NONE-07。"""
        path, _ = self.book('aug_notitle.xlsx', title=None, years=AUG_YEARS,
                            heads=AUG_HEADS, merges=AUG_MERGES)
        with self.assertRaises(RuntimeError) as cm:
            parse(path)
        msg = str(cm.exception)
        for token in ('aug_notitle.xlsx', 'Monthly Metrics', 'D..I', 'K6=2026'):
            self.assertIn(token, msg)
        self.assertNotIn('NONE', msg)

    def test_title_disagrees_with_inference_raises(self):
        """标题写 September 2026，推算出的最后一列是 2026-08 ⇒ 两个锚点对不上，抛。"""
        path, _ = self.book('aug_badtitle.xlsx', title='Monthly Metrics Report for September 2026',
                            years=AUG_YEARS, heads=AUG_HEADS, merges=AUG_MERGES)
        with self.assertRaises(RuntimeError) as cm:
            parse(path)
        self.assertIn('2026-09', str(cm.exception))
        self.assertIn('2026-08', str(cm.exception))

    def test_wrong_surviving_year_label_caught_by_title(self):
        """剩下的那个年份格自己填错（K6=2025）⇒ 整排推成 2024-07..2025-08。
        这正是「只凭一个年份格去推」会出的错；标题 August 2026 对不上，必须抛。"""
        path, _ = self.book('aug_badyear.xlsx', title=AUG_TITLE, years={'K': 2025},
                            heads=AUG_HEADS, merges=AUG_MERGES)
        with self.assertRaises(RuntimeError) as cm:
            parse(path)
        self.assertIn('2025-08', str(cm.exception))


class TestControlGroup(_Books):
    """对照组：修复前后都该绿。"""

    def test_july_layout_unchanged_and_silent(self):
        """7 月版式（D6=2025、L6=2026 都填了）⇒ 2025-06..2026-07，且一个字都不打。"""
        path, rows = self.book('jul.xlsx', title=JUL_TITLE, years=JUL_YEARS,
                               heads=JUL_HEADS, merges=JUL_MERGES)
        m, log = parse(path)
        self.assertEqual([str(p) for p in m.index], months('2025-06', '2026-07'))
        self.assertEqual(m.loc[pd.Period('2026-07', 'M'), 'net_deposits_usdbn'],
                         rows['net_deposits_usdbn'] * 100 + 13)
        self.assertEqual(log, '')

    def test_stale_title_does_not_block_complete_year_labels(self):
        """年份格齐全、标题却没改（还写 June 2026）⇒ 照常读。标题只在替上游补年份时才拿来互证，
        不能让「作者忘改标题」这种不影响数据的疏忽把整月拦下 —— 那是本次 bug 的同一种形状。"""
        path, _ = self.book('jul_stale.xlsx', title='Monthly Metrics Report for June 2026',
                            years=JUL_YEARS, heads=JUL_HEADS, merges=JUL_MERGES)
        m, _log = parse(path)
        self.assertEqual([str(p) for p in m.index], months('2025-06', '2026-07'))


class TestLoudFailures(_Books):
    """旧写法下读错也不出声的坏法。"""

    def test_no_year_label_at_all_raises(self):
        """第 6 行一个年份格都没有 ⇒ 抛。标题只互证，不单独定年份。"""
        path, _ = self.book('noyear.xlsx', title=AUG_TITLE, years={},
                            heads=AUG_HEADS, merges=AUG_MERGES)
        with self.assertRaisesRegex(RuntimeError, '一个年份格都没有'):
            parse(path)

    def test_gap_in_month_header_raises(self):
        """年份格齐全，但删掉 F7「Sep」⇒ Aug 后面直接是 Oct。
        旧写法照读不误，页面上 2025-09 这个月就悄悄没了。"""
        heads = dict(AUG_HEADS)
        del heads['F']
        path, _ = self.book('gap.xlsx', title=AUG_TITLE, years={'D': 2025, 'K': 2026}, heads=heads)
        with self.assertRaises(RuntimeError) as cm:
            parse(path)
        self.assertIn('E7「Aug」', str(cm.exception))
        self.assertIn('G7「Oct」', str(cm.exception))

    def test_contradicting_year_labels_raise(self):
        """D6=2025、K6=2027 ⇒ Dec 2025 后面紧跟 Jan 2027，不可能同时成立。
        旧写法照读，把 2026 年的 8 个月整排写成 2027 年。"""
        path, _ = self.book('contra.xlsx', title=AUG_TITLE, years={'D': 2025, 'K': 2027},
                            heads=AUG_HEADS, merges=AUG_MERGES)
        with self.assertRaisesRegex(RuntimeError, '互相矛盾'):
            parse(path)

    def test_quarterly_sheet_cross_checks_with_quarter_of_report_month(self):
        """季度表自己没有标题：Q3 Q4 | Q1 Q2 只有 Q1 头上有 2026 ⇒ 借月度表报告月所在季度互证。
        报告月 2026-06（2026Q2）对得上就推成 2025Q3..2026Q2；报告月 2026-09（2026Q3）就抛。"""
        W = 8
        rows = [[None] * W for _ in range(6)]
        rows[2][5] = 2026                                       # F3
        for c, q in ((2, 'Q3'), (3, 'Q4'), (5, 'Q1'), (6, 'Q2')):  # C4 D4 | F4 G4，E 列空隔
            rows[3][c] = q
        with contextlib.redirect_stdout(io.StringIO()):
            got = hood._periods(rows, W, 'Q', 'x.xlsx「Quarterly KPIs」', pd.Period('2026-06', 'M'))
        self.assertEqual([str(p) for _i, p in got], ['2025Q3', '2025Q4', '2026Q1', '2026Q2'])
        with self.assertRaises(RuntimeError):
            hood._periods(rows, W, 'Q', 'x.xlsx「Quarterly KPIs」', pd.Period('2026-09', 'M'))


@unittest.skipUnless(os.path.exists(REAL_AUG),
                     'cache/ 里没有 8 月那份工作簿（cache 不进 git）—— skip 不等于通过')
class TestRealCachedWorkbook(unittest.TestCase):

    def test_aug_2026_real_file(self):
        """09-10 挂出的真文件：表头 2025-07..2026-08；其中 Jul-26 一列与已入库的 2026-07 行
        25 格逐格相同。错一列的话这 25 格不可能全等（Jun-26 与 Jul-26 的 funded customers 就差 0.1）。"""
        m, log = parse(REAL_AUG)
        self.assertEqual([str(p) for p in m.index], months('2025-07', '2026-08'))
        self.assertIn('D..I', log)
        old = pd.read_csv(os.path.join(os.path.dirname(HERE), 'series', 'hood.csv'),
                          dtype=str).set_index('month')
        for c in old.columns:
            self.assertAlmostEqual(float(m.loc[pd.Period('2026-07', 'M'), c]),
                                   float(old.loc['2026-07', c]), places=9, msg=c)


if __name__ == '__main__':
    unittest.main(verbosity=2)
